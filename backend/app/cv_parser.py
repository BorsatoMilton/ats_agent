"""
Extraccion de texto plano de CVs en PDF o DOCX.
"""
import io

from pypdf import PdfReader
from docx import Document


class FormatoNoSoportado(Exception):
    pass


def extraer_texto(nombre_archivo: str, contenido: bytes) -> str:
    extension = nombre_archivo.lower().rsplit(".", 1)[-1] if "." in nombre_archivo else ""

    if extension == "pdf":
        return _extraer_de_pdf(contenido)
    if extension == "docx":
        return _extraer_de_docx(contenido)

    raise FormatoNoSoportado(
        f"Formato '.{extension}' no soportado. Se acepta .pdf y .docx"
    )


def _extraer_de_pdf(contenido: bytes) -> str:
    reader = PdfReader(io.BytesIO(contenido))
    partes = []
    for page in reader.pages:
        texto = page.extract_text() or ""
        if texto.strip():
            partes.append(texto)
    return "\n\n".join(partes).strip()


def _extraer_de_docx(contenido: bytes) -> str:
    doc = Document(io.BytesIO(contenido))
    partes = [p.text for p in doc.paragraphs if p.text.strip()]

    # Tablas del docx tambien pueden tener info relevante (ej: skills en
    # una tabla de dos columnas)
    for tabla in doc.tables:
        for fila in tabla.rows:
            celdas = [c.text.strip() for c in fila.cells if c.text.strip()]
            if celdas:
                partes.append(" | ".join(celdas))

    return "\n".join(partes).strip()
